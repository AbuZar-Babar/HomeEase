import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

SOURCE_PPTX = r"C:\Users\AbuZar\Desktop\Fyp\HomeEase\Docs\references\TrekPal_Benchmark\Trekpal FYP 60%.pptx"
DEST_PPTX = r"C:\Users\AbuZar\Desktop\Fyp\HomeEase\Docs\03_Thesis_60\HomeEase_Presentation_60.pptx"
DEST_PPTX_ROOT = r"C:\Users\AbuZar\Desktop\Fyp\HomeEase\Docs\HomeEase_Presentation_60.pptx"

DIAGRAMS_DIR = r"C:\Users\AbuZar\Desktop\Fyp\HomeEase\Docs\03_Thesis_60\diagrams"
SCHEDULE_IMG = os.path.join(DIAGRAMS_DIR, "Schedule_Gantt.png")

# Copy source presentation to work on a fresh clone with exact master and layouts
shutil.copyfile(SOURCE_PPTX, DEST_PPTX)
prs = Presentation(DEST_PPTX)

print(f"Loaded {len(prs.slides)} slides from template.")

def set_font_run(run, font_name="Times New Roman", size_pt=None, bold=None, color_rgb=None):
    run.font.name = font_name
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if color_rgb is not None:
        run.font.color.rgb = color_rgb

def remove_shape(shape):
    sp_elem = shape._element
    sp_elem.getparent().remove(sp_elem)

def get_slide_number_shape(slide):
    for s in slide.shapes:
        if s.name == "Slide Number Placeholder 3":
            return s
    return None

def set_slide_number(slide, num):
    sn = get_slide_number_shape(slide)
    if sn and sn.has_text_frame:
        sn.text_frame.text = str(num)
        for p in sn.text_frame.paragraphs:
            for r in p.runs:
                set_font_run(r, "Times New Roman", 12)

def fit_picture_centered(slide, img_path, left_bound, top_bound, max_w, max_h):
    im = Image.open(img_path)
    im_w, im_h = im.size
    aspect = im_w / im_h

    if max_w / max_h > aspect:
        fitted_h = max_h
        fitted_w = int(max_h * aspect)
    else:
        fitted_w = max_w
        fitted_h = int(max_w / aspect)

    left = left_bound + (max_w - fitted_w) // 2
    top = top_bound + (max_h - fitted_h) // 2

    return slide.shapes.add_picture(img_path, left, top, fitted_w, fitted_h)

def remove_existing_pictures(slide):
    pics = [s for s in slide.shapes if s.shape_type == 13]
    for p in pics:
        remove_shape(p)

# ==============================================================================
# SLIDE 1: Title Slide
# ==============================================================================
print("Configuring Slide 1: Title")
s1 = prs.slides[0]
set_slide_number(s1, 1)

# Title placeholder
title_s1 = s1.shapes.title
title_s1.text = "Project Title:\nHomeEase"
for p in title_s1.text_frame.paragraphs:
    for r in p.runs:
        set_font_run(r, "Times New Roman", 40, bold=True)

# Subtitle 2 (Placeholder)
subtitle_shape = None
for s in s1.shapes:
    if s.name == "Subtitle 2" and s.shape_type == 14: # Placeholder
        subtitle_shape = s
        break

if subtitle_shape:
    tf = subtitle_shape.text_frame
    tf.text = "" # clear
    
    p0 = tf.paragraphs[0]
    p0.text = "Group Members:"
    set_font_run(p0.runs[0], "Times New Roman", 16, bold=True)
    
    members = [
        ("Hashir Hamid", "CIIT/FA22-BSE-139/ATD"),
        ("Aman Ullah Khan", "CIIT/FA22-BSE-074/ATD"),
        ("Umar Saeed", "CIIT/FA22-BCS-041/ATD")
    ]
    for name, roll in members:
        p = tf.add_paragraph()
        p.text = f"{name}\t\t\t({roll})"
        set_font_run(p.runs[0], "Times New Roman", 14)
        
    p_blank = tf.add_paragraph()
    p_blank.text = ""
    
    p_sup = tf.add_paragraph()
    p_sup.text = "Supervisor-name: "
    set_font_run(p_sup.runs[0], "Times New Roman", 16, bold=True)
    
    p_sup_name = tf.add_paragraph()
    p_sup_name.text = "Sir Sumair Khan"
    set_font_run(p_sup_name.runs[0], "Times New Roman", 15)

# CUI Logo
logo_path = os.path.join(DIAGRAMS_DIR, "cui_logo.jpeg")
for s in list(s1.shapes):
    if s.shape_type == 13:
        remove_shape(s)
if os.path.exists(logo_path):
    s1.shapes.add_picture(logo_path, 152473, 451485, 996315, 996315)

# Subtitle 2 Textbox (Department & Campus)
dept_shape = None
for s in s1.shapes:
    if s.name == "Subtitle 2" and s.shape_type == 17: # Text box
        dept_shape = s
        break
if dept_shape:
    tf = dept_shape.text_frame
    tf.text = "Department of Computer Science\nCOMSATS University Islamabad, Abbottabad Campus"
    p0 = tf.paragraphs[0]
    p1 = tf.paragraphs[1]
    set_font_run(p0.runs[0], "Times New Roman", 20, bold=True)
    set_font_run(p1.runs[0], "Times New Roman", 16)


# ==============================================================================
# SLIDE 2: Agenda
# ==============================================================================
print("Configuring Slide 2: Agenda")
s2 = prs.slides[1]
set_slide_number(s2, 2)
s2.shapes.title.text = "Agenda of the Presentation"
for r in s2.shapes.title.text_frame.paragraphs[0].runs:
    set_font_run(r, "Times New Roman", 36)

agenda_items = [
    "Brief Introduction",
    "Scope",
    "Functional Requirements",
    "Non-Functional Requirements",
    "User Goals",
    "Use Cases",
    "Use Case Diagram",
    "Architecture Diagram",
    "Process Flow Diagram",
    "Class Diagram",
    "Sequence Diagram",
    "Activity/State Machine Diagram",
    "ER Diagram",
    "Summary of Project Implementation (60% Milestone)",
    "Schedule (Gantt Chart)",
    "Details of Next Iteration (100% Roadmap)",
    "Conclusion"
]
cp2 = None
for s in s2.shapes:
    if s.name == "Content Placeholder 2":
        cp2 = s
        break
if cp2:
    tf = cp2.text_frame
    tf.text = ""
    for idx, item in enumerate(agenda_items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        set_font_run(p.runs[0], "Times New Roman", 15)


# ==============================================================================
# SLIDE 3: Brief Introduction
# ==============================================================================
print("Configuring Slide 3: Brief Introduction")
s3 = prs.slides[2]
set_slide_number(s3, 3)
s3.shapes.title.text = "Brief Introduction"
for r in s3.shapes.title.text_frame.paragraphs[0].runs:
    set_font_run(r, "Times New Roman", 36)

intro_texts = [
    ("HomeEase", " is an enterprise-grade, hyper-local domestic worker connect platform and marketplace engineered to formalize, organize, and secure the informal home services sector in Pakistan (piloted in Abbottabad)."),
    ("The platform bridges the deep trust deficit between urban households needing domestic help (cooks, cleaners, maids, child nannies, elderly caregivers) and informal workers seeking reliable employment, fair wages, and reputation portability.", ""),
    ("It addresses the critical challenges of informal hiring, such as absence of background checks, safety vulnerabilities, arbitrary wages, and unpredictable worker absenteeism.", ""),
    ("A key innovation of HomeEase is its ", "Content-Based AI Recommender Engine with Explainable AI (XAI) transparent match badges combining Cosine Similarity on skill vectors, Haversine geospatial proximity decay, and min-max rating normalization."),
    ("The platform introduces a ", "Bidirectional Job Marketplace allowing households to publish open gigs and domestic workers to actively browse and apply for local neighbourhood opportunities."),
    ("The system comprises a ", "Flutter mobile app, Supabase PostgreSQL backend with Row Level Security (RLS), and native bilingual English/Urdu (اردو) localization with low-literacy pictorial affordances.")
]

target_s3 = None
for s in s3.shapes:
    if s.has_text_frame and s != s3.shapes.title and s.name != "Slide Number Placeholder 3":
        target_s3 = s
        break
if target_s3:
    tf = target_s3.text_frame
    tf.text = ""
    for idx, (t1, t2) in enumerate(intro_texts):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if t2:
            r1 = p.add_run()
            r1.text = t1
            set_font_run(r1, "Times New Roman", 15, bold=True)
            r2 = p.add_run()
            r2.text = t2
            set_font_run(r2, "Times New Roman", 15)
        else:
            r1 = p.add_run()
            r1.text = t1
            set_font_run(r1, "Times New Roman", 15)


# ==============================================================================
# SLIDE 4: Scope
# ==============================================================================
print("Configuring Slide 4: Scope")
s4 = prs.slides[3]
set_slide_number(s4, 4)
s4.shapes.title.text = "Scope"
for r in s4.shapes.title.text_frame.paragraphs[0].runs:
    set_font_run(r, "Times New Roman", 36)

scope_texts_1 = [
    "Provide a cross-platform Flutter mobile application for households to discover verified domestic workers, inspect XAI match scores, create direct bookings, and post open gigs.",
    "Provide a dedicated worker mobile interface with high-contrast pictorial icons, Urdu voice/text affordances, availability schedule management, and neighborhood gig feeds.",
    "Implement an authentic Content-Based AI Recommender System matching household chore requirements to worker trade skill vectors and Abbottabad geospatial coordinates.",
    "Deliver a Bidirectional Job Marketplace empowering domestic workers to proactively browse and apply to open gigs posted by nearby households.",
    "Auto-generate formal digital service agreements and provide transparent manual payment logging (Cash, JazzCash, EasyPaisa) with audit trails."
]
target_s4 = None
for s in s4.shapes:
    if s.has_text_frame and s != s4.shapes.title and s.name != "Slide Number Placeholder 3":
        target_s4 = s
        break
if target_s4:
    tf = target_s4.text_frame
    tf.text = ""
    for idx, item in enumerate(scope_texts_1):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        set_font_run(p.runs[0], "Times New Roman", 16)


# ==============================================================================
# SLIDE 5: Scope(cont)
# ==============================================================================
print("Configuring Slide 5: Scope(cont)")
s5 = prs.slides[4]
set_slide_number(s5, 5)
s5.shapes.title.text = "Scope(cont)"
for r in s5.shapes.title.text_frame.paragraphs[0].runs:
    set_font_run(r, "Times New Roman", 36)

scope_texts_2 = [
    "Provide a React/Vite web administrative dashboard for KYC document inspection, CNIC verification, dispute mediation, and system analytics.",
    "Enforce database-level scheduling integrity and double-booking locks preventing conflicting appointments for domestic workers.",
    "Implement multi-dimensional review mechanisms evaluating punctuality, professional behavior, and task completion quality.",
    "Implement Supabase JWT authentication and 24 Row Level Security (RLS) policies guaranteeing multi-tenant data isolation and privacy.",
    "Validate system performance with an Abbottabad cold-start evaluation dataset covering real local neighborhoods (Mandian, Jinnahabad, Supply, PMA Kakul, Murree Road)."
]
target_s5 = None
for s in s5.shapes:
    if s.has_text_frame and s != s5.shapes.title and s.name != "Slide Number Placeholder 3":
        target_s5 = s
        break
if target_s5:
    tf = target_s5.text_frame
    tf.text = ""
    for idx, item in enumerate(scope_texts_2):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        set_font_run(p.runs[0], "Times New Roman", 16)


# ==============================================================================
# Helper for Functional Requirements Slides (6, 7, 8, 9)
# ==============================================================================
def populate_fr_slide(slide, slide_num, title_text, section_header, fr_items):
    set_slide_number(slide, slide_num)
    slide.shapes.title.text = title_text
    for r in slide.shapes.title.text_frame.paragraphs[0].runs:
        set_font_run(r, "Times New Roman", 36)
        
    target = None
    for s in slide.shapes:
        if s.has_text_frame and s != slide.shapes.title and s.name != "Slide Number Placeholder 3":
            target = s
            break
    if target:
        tf = target.text_frame
        tf.text = ""
        p_head = tf.paragraphs[0]
        p_head.text = section_header
        set_font_run(p_head.runs[0], "Times New Roman", 20, bold=True)
        
        for name, desc in fr_items:
            p = tf.add_paragraph()
            r1 = p.add_run()
            r1.text = name + ": "
            set_font_run(r1, "Times New Roman", 14, bold=True)
            r2 = p.add_run()
            r2.text = desc
            set_font_run(r2, "Times New Roman", 14)

# SLIDE 6: FR - Household
print("Configuring Slide 6: FR Household")
fr_hh = [
    ("Household Registration & Profile", "Allows households to create an account, manage contact information, and configure Abbottabad locality."),
    ("AI Worker Discovery & Search", "Enables households to filter domestic workers by category, budget, and distance, receiving real-time Content-Based AI rankings."),
    ("Explainable AI Match Badging", "Allows households to view transparent match percentages detailing skill alignment, distance, and past ratings."),
    ("Direct Booking Request Formulation", "Enables households to schedule direct service requests specifying start dates, hours, and chore checklists."),
    ("Publish Open Neighborhood Gigs", "Allows households to post customized chore requirements with defined budget and timeline to the job board."),
    ("Review Worker Applications", "Enables households to evaluate applying domestic workers for posted gigs and accept the best candidate."),
    ("Digital Service Agreement Sign-off", "Allows households to review and digitally sign auto-generated employment contracts prior to job commencement."),
    ("Manual Payment Logging & Receipts", "Allows households to record payments via Cash, JazzCash, or EasyPaisa and attach screenshot receipts.")
]
populate_fr_slide(prs.slides[5], 6, "Functional Requirements", "Household Functional Requirements", fr_hh)

# SLIDE 7: FR - Domestic Worker
print("Configuring Slide 7: FR Worker")
fr_wk = [
    ("Worker Registration & Profile Management", "Allows workers to create a profile, specify bio, years of experience, and base Abbottabad neighborhood."),
    ("KYC & CNIC Document Submission", "Enables workers to upload high-resolution front/back CNIC images and reference certificates for administrative vetting."),
    ("Service Category & Pricing Configuration", "Allows workers to select service specialties (Cooking, Cleaning, Childcare, Elderly Care, Maid) and configure hourly or visit rates."),
    ("Weekly Availability Slot Management", "Provides a weekly scheduling calendar enabling workers to toggle availability time slots by day."),
    ("Bidirectional Job Board Feed", "Displays open gigs posted by nearby households matching registered trades and geographical radius."),
    ("One-Tap Job Application Submission", "Enables domestic workers to proactively apply for open gigs with proposed rates and notes."),
    ("Booking Invitation Management", "Allows workers to inspect, accept, or decline incoming household booking requests in real-time."),
    ("Payment Receipt Confirmation", "Allows workers to inspect household payment logs and confirm receipt of cash or mobile wallet funds."),
    ("Bilingual Urdu Navigation", "Enables workers to access complete app features in Urdu (اردو) supported by high-contrast pictorial visual icons.")
]
populate_fr_slide(prs.slides[6], 7, "Functional Requirements(Cont.)", "Domestic Worker Functional Requirements", fr_wk)

# SLIDE 8: FR - AI Recommendation & Marketplace
print("Configuring Slide 8: FR AI & Marketplace")
fr_ai = [
    ("Skill Vectorization", "Transforms household chore demands and worker capabilities into multi-hot binary attribute vectors across 50 skills."),
    ("Cosine Similarity Computation", "Computes mathematical dot-product similarity between worker skill vectors and household demands."),
    ("Haversine Proximity Decay", "Calculates geodesic distance in kilometers between coordinates and applies exponential spatial decay penalties."),
    ("Composite Match Scoring", "Computes weighted score combining skill similarity (50%), proximity (30%), and normalized ratings (20%)."),
    ("Explainable AI (XAI) Badging", "Generates human-readable badges explaining why the worker was recommended (skills, distance, ratings)."),
    ("Bidirectional Feed Matching", "Filters open gigs for domestic workers matching their approved categories within their travel radius."),
    ("Double-Booking Conflict Lock", "Enforces transactional database constraints preventing overlapping bookings for the same worker."),
    ("Dispute Ticket Logging", "Allows either party to log formal dispute tickets with photographic evidence and detailed claims.")
]
populate_fr_slide(prs.slides[7], 8, "Functional Requirements(Cont.)", "AI Recommendation & Marketplace Functional Requirements", fr_ai)

# SLIDE 9: FR - Admin & Governance
print("Configuring Slide 9: FR Admin")
fr_adm = [
    ("Administrator Authentication", "Provides secure role-based dashboard access for university evaluators and platform managers."),
    ("Worker Verification Queue", "Allows administrators to inspect submitted CNIC cards, verify identity records, and approve or reject workers."),
    ("Dispute Mediation Console", "Enables administrators to inspect disputed bookings, review photographic evidence and payment logs, and issue binding resolutions."),
    ("User & Booking Audit Ledgers", "Allows administrators to monitor global platform activity, track booking lifecycles, and inspect audit records."),
    ("Platform Taxonomy & Skill Management", "Enables administrators to manage domestic service categories, skills, and standardized rates."),
    ("Data Export & Compliance Reporting", "Supports exporting platform activity, user rosters, and dispute statistics in CSV format for academic evaluation.")
]
populate_fr_slide(prs.slides[8], 9, "Functional Requirements(Cont.)", "Admin Functional Requirements", fr_adm)


# ==============================================================================
# SLIDE 10: Non-Functional Requirements
# ==============================================================================
print("Configuring Slide 10: NFR")
s10 = prs.slides[9]
set_slide_number(s10, 10)
s10.shapes.title.text = "Non-Functional Requirements"
for r in s10.shapes.title.text_frame.paragraphs[0].runs:
    set_font_run(r, "Times New Roman", 36)

nfr_sections = [
    ("Usability Requirements", [
        ("NFR-USE-01 — Low-Literacy Accessibility", "All critical worker actions (accept job, view feed, confirm pay) executable within ≤3 taps from home screen. High-contrast pictorial iconography for all tasks."),
        ("NFR-USE-02 — Native Urdu & RTL", "Flawless bilingual layout switching with correct Right-to-Left (RTL) padding and Urdu typography.")
    ]),
    ("Reliability Requirements", [
        ("NFR-REL-01 — High Availability", "The system shall maintain 99.5% uptime backed by Supabase cloud infrastructure."),
        ("NFR-REL-02 — Prevent Double Booking", "The system shall ensure no worker can be double-booked across overlapping calendar time slots.")
    ]),
    ("Performance Requirements", [
        ("NFR-PERF-01 — Fast Response", "Content-Based AI recommendation engine computes match scores across 500 worker profiles in <250 ms."),
        ("NFR-PERF-02 — Scalable User Load", "Mobile UI maintains 60 FPS transitions with screen loading times under 2 seconds.")
    ])
]

target_s10 = None
for s in s10.shapes:
    if s.has_text_frame and s != s10.shapes.title and s.name != "Slide Number Placeholder 3":
        target_s10 = s
        break
if target_s10:
    tf = target_s10.text_frame
    tf.text = ""
    first_p = True
    for cat_title, items in nfr_sections:
        p_cat = tf.paragraphs[0] if first_p else tf.add_paragraph()
        first_p = False
        p_cat.text = cat_title
        set_font_run(p_cat.runs[0], "Times New Roman", 18, bold=True)
        
        for code_name, desc in items:
            p = tf.add_paragraph()
            r1 = p.add_run()
            r1.text = code_name + "\n"
            set_font_run(r1, "Times New Roman", 13, bold=True)
            r2 = p.add_run()
            r2.text = desc
            set_font_run(r2, "Times New Roman", 13)


# ==============================================================================
# SLIDE 11: User Goals
# ==============================================================================
print("Configuring Slide 11: User Goals")
s11 = prs.slides[10]
set_slide_number(s11, 11)
s11.shapes.title.text = "User Goals"
for r in s11.shapes.title.text_frame.paragraphs[0].runs:
    set_font_run(r, "Times New Roman", 36)

user_goals = [
    ("Find a ", "simple and centralized way", " to discover and hire verified domestic help with complete safety."),
    ("Evaluate ", "transparent Explainable AI match scores", " combining skills, distance, and authentic community ratings."),
    ("Empower domestic workers with ", "active marketplace agency", " to browse open household jobs and apply proactively."),
    ("Prevent disputes and wage ambiguity through ", "formal digital service agreements", " and manual payment audit logs."),
    ("Overcome language and literacy barriers through ", "native Urdu localization and visual pictorial guides", "."),
    ("Save time and eliminate stress by replacing ", "unreliable word-of-mouth networks", " with a single, trusted mobile platform.")
]

target_s11 = None
for s in s11.shapes:
    if s.has_text_frame and s != s11.shapes.title and s.name != "Slide Number Placeholder 3":
        target_s11 = s
        break
if target_s11:
    tf = target_s11.text_frame
    tf.text = ""
    for idx, (pfx, bold_txt, sfx) in enumerate(user_goals):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        r1 = p.add_run()
        r1.text = pfx
        set_font_run(r1, "Times New Roman", 17)
        r2 = p.add_run()
        r2.text = bold_txt
        set_font_run(r2, "Times New Roman", 17, bold=True)
        r3 = p.add_run()
        r3.text = sfx
        set_font_run(r3, "Times New Roman", 17)


# ==============================================================================
# Helper for Use Cases Slides (12, 13, 14, 15)
# ==============================================================================
def populate_uc_slide(slide, slide_num, title_text, uc_category, uc_items):
    set_slide_number(slide, slide_num)
    slide.shapes.title.text = title_text
    for r in slide.shapes.title.text_frame.paragraphs[0].runs:
        set_font_run(r, "Times New Roman", 36)
        
    target = None
    for s in slide.shapes:
        if s.has_text_frame and s != slide.shapes.title and s.name != "Slide Number Placeholder 3":
            target = s
            break
    if target:
        tf = target.text_frame
        tf.text = ""
        p_head = tf.paragraphs[0]
        p_head.text = uc_category
        set_font_run(p_head.runs[0], "Times New Roman", 22, bold=True)
        
        for uc in uc_items:
            p = tf.add_paragraph()
            p.text = uc
            set_font_run(p.runs[0], "Times New Roman", 16)

# SLIDE 12: Household Use Cases
print("Configuring Slide 12: Household Use Cases")
uc_hh = [
    "UC-1  Register / Login",
    "UC-2  Profile & Abbottabad Locality Setup",
    "UC-3  Discover Workers via AI Recommender",
    "UC-4  View Explainable AI (XAI) Match Badges",
    "UC-5  Send Direct Booking Request",
    "UC-6  Publish Open Household Gig",
    "UC-7  Review Worker Job Applications",
    "UC-8  Sign Digital Service Agreement",
    "UC-9  Log Payment & Upload Receipt",
    "UC-10 Submit Multi-Dimensional Review"
]
populate_uc_slide(prs.slides[11], 12, "Use Cases", "Household Use Cases", uc_hh)

# SLIDE 13: Domestic Worker Use Cases
print("Configuring Slide 13: Domestic Worker Use Cases")
uc_wk = [
    "UC-11 Register / Login",
    "UC-12 Setup Trade Profile & Service Categories",
    "UC-13 Submit KYC Verification (CNIC Upload)",
    "UC-14 Configure Weekly Availability Calendar",
    "UC-15 Browse Open Neighborhood Gigs Feed",
    "UC-16 Submit One-Tap Job Application",
    "UC-17 View & Accept Direct Booking Invitations",
    "UC-18 Sign Service Agreement & Execute Chores",
    "UC-19 Confirm Payment Receipt"
]
populate_uc_slide(prs.slides[12], 13, "Use Cases(cont)", "Domestic Worker Use Cases", uc_wk)

# SLIDE 14: Admin Use Cases
print("Configuring Slide 14: Admin Use Cases")
uc_adm = [
    "UC-20 Admin Authentication & Console Login",
    "UC-21 Review & Approve/Reject Worker KYC (CNIC)",
    "UC-22 Mediate & Resolve Booking Disputes",
    "UC-23 View System Analytics & User Rosters",
    "UC-24 Manage Domestic Service Categories & Skills",
    "UC-25 Suspend / Deactivate Non-Compliant Users",
    "UC-26 Export Platform Audit Logs (CSV)"
]
populate_uc_slide(prs.slides[13], 14, "Use Cases(cont)", "Admin Use Cases", uc_adm)

# SLIDE 15: System & Platform Use Cases
print("Configuring Slide 15: System & Platform Use Cases")
uc_sys = [
    "UC-27 Compute Content-Based Cosine Similarity",
    "UC-28 Compute Haversine Geospatial Proximity Decay",
    "UC-29 Generate Transparent Explainable AI (XAI) Badges",
    "UC-30 Filter Bidirectional Job Feed by Proximity",
    "UC-31 Enforce Double-Booking Concurrency Lock",
    "UC-32 Dynamic Bilingual Language Toggle (EN / اردو)",
    "UC-33 Render High-Affordance Pictorial Task Cards"
]
populate_uc_slide(prs.slides[14], 15, "Use Cases(cont)", "System & Platform Use Cases", uc_sys)


# ==============================================================================
# Helper for Diagram Slides (16, 17, 18, 19, 20, 21, 22, 25)
# ==============================================================================
def populate_diagram_slide(slide, slide_num, title_text, img_path, left_bound=457200, top_bound=1200000, max_w=8229600, max_h=5100000):
    set_slide_number(slide, slide_num)
    slide.shapes.title.text = title_text
    for r in slide.shapes.title.text_frame.paragraphs[0].runs:
        set_font_run(r, "Times New Roman", 36)
        
    remove_existing_pictures(slide)
    
    if os.path.exists(img_path):
        fit_picture_centered(slide, img_path, left_bound, top_bound, max_w, max_h)
    else:
        print(f"WARNING: Image not found: {img_path}")

# SLIDE 16: Use Case Diagram
print("Configuring Slide 16: Use Case Diagram")
populate_diagram_slide(prs.slides[15], 16, "Use Case Diagram", os.path.join(DIAGRAMS_DIR, "Use_Case_Diagram.png"))

# SLIDE 17: Architecture Diagram
print("Configuring Slide 17: Architecture Diagram")
populate_diagram_slide(prs.slides[16], 17, "Architecture Diagram", os.path.join(DIAGRAMS_DIR, "Architecture.png"), left_bound=457200, top_bound=1250000, max_w=8229600, max_h=5100000)

# SLIDE 18: Process Flow Diagram
print("Configuring Slide 18: Process Flow Diagram")
populate_diagram_slide(prs.slides[17], 18, "Process Flow Diagram", os.path.join(DIAGRAMS_DIR, "ProcessFlow.png"), left_bound=457200, top_bound=1200000, max_w=8229600, max_h=5200000)

# SLIDE 19: Class Diagram
print("Configuring Slide 19: Class Diagram")
populate_diagram_slide(prs.slides[18], 19, "Class Diagram", os.path.join(DIAGRAMS_DIR, "Class.png"))

# SLIDE 20: Sequence Diagram
print("Configuring Slide 20: Sequence Diagram")
populate_diagram_slide(prs.slides[19], 20, "Sequence Diagram", os.path.join(DIAGRAMS_DIR, "Sequence.png"))

# SLIDE 21: Activity/State Machine Diagram
print("Configuring Slide 21: Activity/State Machine Diagram")
populate_diagram_slide(prs.slides[20], 21, "Activity/State Machine Diagram", os.path.join(DIAGRAMS_DIR, "StateTransition.png"))

# SLIDE 22: ER Diagram
print("Configuring Slide 22: ER Diagram")
populate_diagram_slide(prs.slides[21], 22, "ER Diagram", os.path.join(DIAGRAMS_DIR, "ERD.png"))


# ==============================================================================
# Helper for Summary of Implementation Slides (23, 24)
# ==============================================================================
def populate_summary_slide(slide, slide_num, title_text, items):
    set_slide_number(slide, slide_num)
    slide.shapes.title.text = title_text
    for r in slide.shapes.title.text_frame.paragraphs[0].runs:
        set_font_run(r, "Times New Roman", 36)
        
    target = None
    for s in slide.shapes:
        if s.has_text_frame and s != slide.shapes.title and s.name != "Slide Number Placeholder 3":
            target = s
            break
    if target:
        tf = target.text_frame
        tf.text = ""
        for idx, (title, desc) in enumerate(items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            r1 = p.add_run()
            r1.text = title + ": "
            set_font_run(r1, "Times New Roman", 15, bold=True)
            r2 = p.add_run()
            r2.text = desc
            set_font_run(r2, "Times New Roman", 15)

# SLIDE 23: Summary of Project Implementation
print("Configuring Slide 23: Summary of Implementation Part 1")
sum_items_1 = [
    ("Multi-role platform architecture", "HomeEase is designed as a unified ecosystem with dedicated interfaces for households, domestic workers, and administrators, so each user type can perform tasks relevant to their role."),
    ("Cross-platform mobile application", "Built with Flutter & Dart, supporting rapid responsive performance, seamless role-switching, and full onboarding for both households and domestic workers."),
    ("Content-Based AI Recommender System", "Instead of static dropdown filters, the system computes Cosine Similarity on skill vectors, Haversine geospatial proximity decay across Abbottabad coordinates, and min-max rating normalization with transparent Explainable AI (XAI) badges."),
    ("Bidirectional job marketplace", "Allows households to publish open gigs specifying budget, schedule, and chore checklists, while domestic workers actively browse local neighbourhood feeds and apply with single-tap ease."),
    ("Booking & digital service agreements", "End-to-end booking management supporting scheduling, transactional double-booking prevention, auto-generated service agreements with digital sign-off, and structured state transitions.")
]
populate_summary_slide(prs.slides[22], 23, "Summary of Project Implementation", sum_items_1)

# SLIDE 24: Summary of Project Implementation (cont.)
print("Configuring Slide 24: Summary of Implementation Part 2")
sum_items_2 = [
    ("Admin control and verification", "The React/Vite admin dashboard supports inspection and verification of worker CNIC documents, approval/rejection workflows, and platform-level governance ensuring trust across Abbottabad."),
    ("Bilingual localization & visual affordances", "The platform includes complete English and Urdu (اردو) localization with right-to-left layout adaptation and high-contrast pictorial icons for low-literacy informal workers."),
    ("Manual payment audit trail", "Supports logging of Cash, JazzCash, and EasyPaisa transactions with digital receipt screenshot uploads and worker receipt confirmation."),
    ("Dispute resolution workflow", "Allows parties to submit evidence-backed dispute tickets with photos and descriptions for administrative arbitration."),
    ("Secure and scalable relational backend", "Centralized Supabase PostgreSQL database featuring 16 normalized relational tables, strict foreign key constraints, and 24 Row Level Security (RLS) policies guaranteeing data isolation and security.")
]
populate_summary_slide(prs.slides[23], 24, "Summary of Project Implementation (cont.)", sum_items_2)


# ==============================================================================
# SLIDE 25: Schedule (Gantt Chart)
# ==============================================================================
print("Configuring Slide 25: Schedule")
populate_diagram_slide(prs.slides[24], 25, "Schedule", SCHEDULE_IMG, left_bound=457200, top_bound=1400000, max_w=8229600, max_h=4800000)


# ==============================================================================
# Helper for 100% Next Iteration Slides (26, 27, 28)
# ==============================================================================
def populate_iteration_slide(slide, slide_num, sections):
    set_slide_number(slide, slide_num)
    slide.shapes.title.text = "Details of next Iteration (100%)"
    for r in slide.shapes.title.text_frame.paragraphs[0].runs:
        set_font_run(r, "Times New Roman", 36)
        
    target = None
    for s in slide.shapes:
        if s.has_text_frame and s != slide.shapes.title and s.name != "Slide Number Placeholder 3":
            target = s
            break
    if target:
        tf = target.text_frame
        tf.text = ""
        first_p = True
        for header, sub_bullets in sections:
            p_head = tf.paragraphs[0] if first_p else tf.add_paragraph()
            first_p = False
            p_head.text = header
            set_font_run(p_head.runs[0], "Times New Roman", 18, bold=True)
            
            for b in sub_bullets:
                p = tf.add_paragraph()
                p.text = b
                set_font_run(p.runs[0], "Times New Roman", 14)

# SLIDE 26: Next Iteration 1
print("Configuring Slide 26: Next Iteration Part 1")
iter_26 = [
    ("Real-Time Communication System", [
        "Bidirectional in-app live chat between households and domestic workers using Supabase Realtime WebSockets",
        "In-app audio voice messaging to empower illiterate and low-literacy domestic workers",
        "Read receipts and message delivery indicators"
    ]),
    ("Automated Push Notification Infrastructure", [
        "Firebase Cloud Messaging (FCM) integration for real-time mobile push notifications",
        "Automated alerts for new job applications, booking responses, payment confirmations, and reminders",
        "In-app notification center for persistent system alerts and updates"
    ])
]
populate_iteration_slide(prs.slides[25], 26, iter_26)

# SLIDE 27: Next Iteration 2
print("Configuring Slide 27: Next Iteration Part 2")
iter_27 = [
    ("Payment Gateway & Escrow Automation", [
        "Direct integration with Pakistani mobile payment APIs (JazzCash, EasyPaisa, 1Link)",
        "Milestone-based digital escrow holding funds until domestic chores are verified as completed",
        "Detailed receipt downloading and monthly household expenditure statements"
    ]),
    ("Advanced Identity & Security Hardening", [
        "Live Pakistani mobile SMS OTP verification to eliminate fraudulent registrations",
        "NADRA-compliant CNIC verification API integration exploration for enhanced trust",
        "Biometric fingerprint / face unlock for mobile worker login"
    ])
]
populate_iteration_slide(prs.slides[26], 27, iter_27)

# SLIDE 28: Next Iteration 3
print("Configuring Slide 28: Next Iteration Part 3")
iter_28 = [
    ("Dispute Resolution & Policy Enforcement", [
        "Tiered penalty scoring and automated temporary account suspension for repeated cancellations",
        "Comprehensive dispute arbitration console with audit trails and settlement logs"
    ]),
    ("Deployment and Production Hardening", [
        "Migration of KYC documents and receipt uploads to secure S3 / Supabase cloud storage buckets",
        "Empirical usability testing and pilot deployment with real households and workers in Abbottabad",
        "Full regression testing, algorithmic latency benchmarking, and performance optimization",
        "Comprehensive final thesis documentation, user manual, and live defense demonstration"
    ])
]
populate_iteration_slide(prs.slides[27], 28, iter_28)

# Save presentations
prs.save(DEST_PPTX)
shutil.copyfile(DEST_PPTX, DEST_PPTX_ROOT)
print(f"Successfully generated HomeEase 60% Presentation at:\n1) {DEST_PPTX}\n2) {DEST_PPTX_ROOT}")
